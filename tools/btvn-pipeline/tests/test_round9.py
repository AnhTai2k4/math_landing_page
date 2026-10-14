import copy,io,os,unittest,datetime
from unittest.mock import patch
from openpyxl import Workbook
from test_pipeline import p,ROOT

def fixture():
    book=Workbook();ws=book.active;ws.title='Lớp giả'
    ws.append(['Lớp','Ngày','Tên bài','Tài liệu','Handout','Video','BTVN','Đáp án'])
    ws.append(['2k10 A',datetime.date(2026,10,15),'Bài giả','HS.pdf','','','BTVN.pdf','GV.pdf'])
    for col,label in [('D','hs'),('G','btvn'),('H','gv')]:ws[f'{col}2'].hyperlink=f'https://pdf.example.invalid/{label}.pdf'
    raw=io.BytesIO();book.save(raw);return raw.getvalue()

def sheet():
    return {'id':'fake-material-sheet-1','format':'xlsx','material_only_workbook':True,'bearer_env':'FAKE_MATERIAL_TOKEN','columns':dict(zip(['class_name','lesson_date','title','document','handout','video','homework','answers'],['Lớp','Ngày','Tên bài','Tài liệu','Handout','Video','BTVN','Đáp án'])),'tabs':[{'name':'Lớp giả','range':'A1:H10'}]}

class Round9Tests(unittest.TestCase):
    def test_export_preserves_hyperlinks_dates_and_deduplicates_request(self):
        s=sheet();config={'sheets':[s,copy.deepcopy(s)]}
        class Transport:
            calls=[]
            def request(self,url,headers=None,limit=None):
                self.calls.append(url);self.limit=limit
                self.assertions=(headers=={'Authorization':'Bearer fake-only'})
                return fixture()
        transport=Transport()
        with patch.dict(os.environ,{'FAKE_MATERIAL_TOKEN':'fake-only'}):rows=p.read_material_rows(config,transport)
        self.assertEqual(len(transport.calls),1);self.assertIn('/drive/v3/files/fake-material-sheet-1/export?mimeType=',transport.calls[0]);self.assertEqual(transport.limit,10*1024*1024);self.assertTrue(transport.assertions)
        self.assertEqual(len(rows),1);self.assertEqual(rows[0]['class_name'],'2K10 A');self.assertEqual(rows[0]['lesson_date'],'2026-10-15');self.assertEqual(rows[0]['homework'],'https://pdf.example.invalid/btvn.pdf');self.assertEqual(rows[0]['answers'],'https://pdf.example.invalid/gv.pdf')

    def test_no_export_without_material_only_confirmation_or_named_range(self):
        class NoNetwork:
            def request(self,*args,**kwargs):raise AssertionError('No network allowed')
        for key,value in [('material_only_workbook',False),('id','bad/id'),('tabs',[{'name':'Lớp giả'}])]:
            s=sheet();s[key]=value
            with patch.dict(os.environ,{'FAKE_MATERIAL_TOKEN':'fake-only'}),self.assertRaises(p.GateError):p.read_material_rows({'sheets':[s]},NoNetwork())

    def test_reject_formula_missing_tab_and_large_range(self):
        s=p.sheet_sources({'sheets':[sheet()]})[0]
        for mutation in [{'name':'Không có'},{'range':'A1:X999'},{'range':'A:K'},{'range':'A1:K99999'}]:
            with self.assertRaises(p.GateError):p.xlsx_rows(fixture(),{**s,**mutation})
        book=Workbook();book.active.title='Lớp giả';book.active['A1']='=HYPERLINK("https://fake.invalid","HS")';raw=io.BytesIO();book.save(raw)
        with self.assertRaises(p.GateError):p.xlsx_rows(raw.getvalue(),s)

    def test_missing_chip_url_not_guessed(self):
        book=Workbook();ws=book.active;ws.title='Lớp giả';ws.append(['Lớp','Ngày','Tên bài','Tài liệu','Handout','Video','BTVN','Đáp án']);ws.append(['2K10 A','2026-10-15','Giả','','','','BTVN.pdf','GV.pdf']);raw=io.BytesIO();book.save(raw)
        with self.assertRaises(p.GateError):p.sheet_rows(raw.getvalue(),p.sheet_sources({'sheets':[sheet()]})[0])

if __name__=='__main__':unittest.main()
